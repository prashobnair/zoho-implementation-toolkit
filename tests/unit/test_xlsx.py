"""XLSX renderer structure tests (TK-CORE-8): open with openpyxl."""

from __future__ import annotations

import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from zohokit.core.context import RunContext
from zohokit.core.findings import Finding, Severity
from zohokit.modules.migration.report import build_report
from zohokit.reports.xlsx import FINDINGS_COLUMNS, SHEET_NAMES, render_xlsx


def _report():  # type: ignore[no-untyped-def]
    now = datetime(2026, 1, 5, tzinfo=UTC)
    findings = [
        Finding.create(
            module="migration",
            code="unknown_target_field",
            severity=Severity.ERROR,
            entity="people",
            entity_id="Nope__s",
            message="Target field is not present.",
            evidence={"target_field": "Nope__s"},
            discriminator="mapping\0Contacts\0Nope__s",
        )
    ]
    from zohokit.modules import Analysis

    return build_report(
        Analysis(findings=tuple(findings), legacy={}, ready=False),
        ctx=RunContext(now=now, mode="offline"),
        inputs_sha256="abc",
    )


def test_workbook_structure() -> None:
    book = load_workbook(filename=io.BytesIO(render_xlsx(_report())))
    assert book.sheetnames == list(SHEET_NAMES)
    summary = book["Summary"]
    assert summary["A1"].value == "zohokit migration report"
    assert summary["B4"].value == "NOT READY"
    # Severity counts plus a SUM totals formula.
    assert [summary.cell(row=row, column=1).value for row in range(7, 12)] == [
        "error",
        "review",
        "warning",
        "info",
        "total",
    ]
    assert summary.cell(row=7, column=2).value == 1
    assert summary.cell(row=11, column=2).value == "=SUM(B7:B10)"
    findings = book["Findings"]
    assert [findings.cell(row=1, column=col).value for col in range(1, 8)] == list(FINDINGS_COLUMNS)
    assert findings.freeze_panes == "A2"
    assert findings.auto_filter.ref == "A1:G2"
    assert findings.cell(row=2, column=4).value == "unknown_target_field"
    signoff = book["Sign-off"]
    assert signoff["A1"].value == "Sign-off"
    assert [signoff.cell(row=row, column=1).value for row in (3, 4, 5)] == [
        "Preparer",
        "Reviewer",
        "Date",
    ]
    assert all(signoff.cell(row=row, column=2).value in (None, "") for row in (3, 4, 5))
