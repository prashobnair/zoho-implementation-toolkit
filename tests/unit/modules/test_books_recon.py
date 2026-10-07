"""Books month-end recon v2 tests (TK-BK-F1..F6, TK-BK-F8, UC-BK-1..3).

Covers the match-policy strategies, tolerance bounds, staged billing,
credit-note netting, tax modes, entity maps with fingerprint-only org
references, FX exactness, status handling, period windowing, row-level
robustness and the Marigold September-2026 answer key.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from hypothesis import given
from hypothesis import strategies as st

from zohokit.core.context import RunContext
from zohokit.core.findings import Severity
from zohokit.modules.books.entities import (
    EntityEntry,
    EntityMapConfigError,
    deal_in_entity,
    load_entity_map,
    org_fingerprint,
    resolve_org_ids,
)
from zohokit.modules.books.fx import FxConfigError, convert, load_fx_rates, lookup_rate
from zohokit.modules.books.period import in_window, record_date
from zohokit.modules.books.policy import (
    MatchPolicy,
    PolicyConfigError,
    load_policy,
)
from zohokit.modules.books.reconcile import analyze_recon, match_key, within_tolerance
from zohokit.reports import render_json

ROOT = Path(__file__).resolve().parent.parent.parent.parent
MARIGOLD = ROOT / "fixtures" / "books" / "marigold"

ORG_IDS = {"in-entity": "555000001", "us-entity": "555000002", "eu-entity": "555000003"}


def _entities() -> dict[str, EntityEntry]:
    return load_entity_map(MARIGOLD / "entity_map.yaml")


def _policy() -> MatchPolicy:
    return load_policy(MARIGOLD / "policy.yaml")


def _marigold_data() -> dict[str, list[dict[str, object]]]:
    import json as _json

    return _json.loads((MARIGOLD / "recon.json").read_text(encoding="utf-8"))


def _analyze(**overrides: object) -> object:

    data = _marigold_data()
    params: dict[str, object] = {
        "deals": data["deals"],
        "invoices": data["invoices"],
        "credit_notes": data["credit_notes"],
        "policy": _policy(),
        "entities": _entities(),
        "org_ids": ORG_IDS,
        "fx_rates": load_fx_rates(MARIGOLD / "fx_rates.csv"),
        "unavailable": ("eu-entity",),
    }
    params.update(overrides)
    return analyze_recon(**params)  # type: ignore[arg-type]


def _codes(analysis: object) -> set[tuple[str, str]]:
    from zohokit.modules import Analysis as _Analysis

    assert isinstance(analysis, _Analysis)
    return {(item.code, item.entity_id) for item in analysis.findings}


# --- policy loading + schema freshness ---------------------------------------


def test_policy_schema_matches_model() -> None:
    on_disk = (ROOT / "schemas" / "books" / "policy.v1.json").read_text(encoding="utf-8")
    assert on_disk == json.dumps(MatchPolicy.model_json_schema(), indent=2, sort_keys=True) + "\n"


def test_policy_loads_with_defaults() -> None:
    policy = _policy()
    assert policy.key_strategy == "crm_deal_id_in_invoice_custom_field"
    assert policy.tolerance.abs == Decimal("1.00")
    assert policy.tolerance.pct == Decimal("0.5")
    assert policy.credit_notes == "net_off"
    assert policy.include_draft is False
    assert policy.period is not None and policy.period.by == "deal_closing_date"
    assert policy.crm_fields.amount == "Amount"
    assert policy.crm_fields.closing_date == "Closing_Date"
    assert policy.crm_fields.currency == "Currency"
    assert policy.crm_fields.customer == "Account_Name"
    assert policy.crm_fields.billing_plan == "Billing_Plan"


def test_default_tolerance_is_exact_match() -> None:
    policy = MatchPolicy()
    assert policy.tolerance.abs is None
    assert policy.tolerance.pct is None


def test_bad_policy_is_a_config_error(tmp_path: Path) -> None:
    bad = tmp_path / "policy.yaml"
    bad.write_text("key_strategy: ouija_board\n", encoding="utf-8")
    with pytest.raises(PolicyConfigError, match=r"policy\.yaml:1:"):
        load_policy(bad)


def test_bad_entity_map_and_fx_are_config_errors(tmp_path: Path) -> None:
    bad_map = tmp_path / "entities.yaml"
    bad_map.write_text("in-entity: {base_currency: XX}\n", encoding="utf-8")
    with pytest.raises(EntityMapConfigError, match=r"entities\.yaml:1:"):
        load_entity_map(bad_map)
    bad_fx = tmp_path / "fx.csv"
    bad_fx.write_text("date,from,to,rate,source\n2026-09-20,USD,INR,abc,x\n", encoding="utf-8")
    with pytest.raises(FxConfigError, match=r"fx\.csv:2:"):
        load_fx_rates(bad_fx)


def test_unknown_org_ref_names_only_the_ref() -> None:
    with pytest.raises(EntityMapConfigError, match="unknown books org ref"):
        resolve_org_ids(_entities(), {"in-entity": "555000001"})


# --- key strategies: one test per strategy ------------------------------------


def _deal(deal_id: str = "d-1") -> dict[str, object]:
    return {
        "id": deal_id,
        "entity": "in-entity",
        "currency": "INR",
        "net_amount": "100.00",
        "customer_name": "Marigold Labs",
        "closing_date": "2026-09-10",
        "reference_number": "REF-1",
        "salesorder_id": "555000901",
    }


def _invoice(invoice_id: str = "i-1") -> dict[str, object]:
    return {
        "id": invoice_id,
        "entity": "in-entity",
        "currency": "INR",
        "net_amount": "100.00",
        "customer_name": "Marigold Labs",
        "status": "sent",
        "date": "2026-09-11",
        "reference_number": "REF-1",
        "salesorder_id": "555000901",
        "custom_fields": [{"label": "CRM Deal ID", "value": "d-1"}],
    }


def test_custom_field_strategy_matches_deal_id() -> None:
    policy = MatchPolicy(key_strategy="crm_deal_id_in_invoice_custom_field")
    assert match_key(policy, _deal(), _invoice()) is True
    other = _invoice()
    other["custom_fields"] = [{"label": "CRM Deal ID", "value": "d-2"}]
    assert match_key(policy, _deal(), other) is False


def test_reference_number_strategy() -> None:
    policy = MatchPolicy(key_strategy="reference_number")
    assert match_key(policy, _deal(), _invoice()) is True
    other = _invoice()
    other["reference_number"] = "REF-9"
    assert match_key(policy, _deal(), other) is False
    # A matching customer name alone never keys (UC-BK-3).
    nameless = _invoice()
    nameless["reference_number"] = ""
    assert match_key(policy, _deal(), nameless) is False


def test_sales_order_link_strategy() -> None:
    policy = MatchPolicy(key_strategy="sales_order_link")
    assert match_key(policy, _deal(), _invoice()) is True
    other = _invoice()
    other["salesorder_id"] = ""
    other["salesorder_number"] = ""
    assert match_key(policy, _deal(), other) is False
    numbered = _invoice()
    numbered["salesorder_id"] = ""
    numbered["salesorder_number"] = "SO-77"
    deal = _deal()
    deal["salesorder_id"] = ""
    deal["salesorder_number"] = "SO-77"
    assert match_key(policy, deal, numbered) is True


# --- tolerance -----------------------------------------------------------------


def test_tolerance_needs_both_bounds() -> None:
    from zohokit.modules.books.policy import Tolerance

    policy = MatchPolicy(tolerance=Tolerance(abs=Decimal("1.00"), pct=Decimal("0.5")))
    assert within_tolerance(Decimal("0.60"), Decimal("100.00"), policy) is False
    assert within_tolerance(Decimal("0.40"), Decimal("100.00"), policy) is True
    assert within_tolerance(Decimal("2.00"), Decimal("100000.00"), policy) is False


def test_tolerance_applies_only_stated_bounds() -> None:
    from zohokit.modules.books.policy import Tolerance

    # Only pct stated: the default abs is NOT silently ANDed in, so a
    # 60.00 diff on a 100.00 net passes a 100% bound on its own.
    pct_only = MatchPolicy(tolerance=Tolerance(pct=Decimal("100")))
    assert within_tolerance(Decimal("60.00"), Decimal("100.00"), pct_only) is True
    assert within_tolerance(Decimal("110.00"), Decimal("100.00"), pct_only) is False
    # Only abs stated: the pct bound is ignored, so 2.00 passes abs 5.00
    # even though it exceeds 0.5% of a 100.00 net.
    abs_only = MatchPolicy(tolerance=Tolerance(abs=Decimal("5.00")))
    assert within_tolerance(Decimal("2.00"), Decimal("100.00"), abs_only) is True
    assert within_tolerance(Decimal("6.00"), Decimal("100.00"), abs_only) is False
    # Neither stated: exact match only.
    exact = MatchPolicy(tolerance=Tolerance())
    assert within_tolerance(Decimal("0"), Decimal("100.00"), exact) is True
    assert within_tolerance(Decimal("0.01"), Decimal("100.00"), exact) is False


def test_single_mismatch_is_outside_tolerance() -> None:
    analysis = _analyze()
    assert ("outside_tolerance", "d-in-03") in _codes(analysis)
    assert ("within_tolerance", "d-in-01") in _codes(analysis)


# --- staged billing (UC-BK-2) ---------------------------------------------------


def test_staged_billing_sums_to_net() -> None:
    analysis = _analyze()
    assert ("within_tolerance", "d-in-02") in _codes(analysis)


def test_multiple_invoices_without_plan_are_duplicates() -> None:
    data = _marigold_data()
    deals = [dict(deal) for deal in data["deals"] if deal["id"] == "d-in-02"]
    deals[0].pop("Billing_Plan")
    analysis = _analyze(deals=deals)
    assert ("duplicate_invoice_reference", "d-in-02") in _codes(analysis)
    assert ("within_tolerance", "d-in-02") not in _codes(analysis)


def test_partial_billing_gives_under_and_over() -> None:
    analysis = _analyze()
    assert ("under_invoiced", "d-in-04") in _codes(analysis)
    assert ("over_invoiced", "d-in-05") in _codes(analysis)


# --- credit notes ---------------------------------------------------------------


def test_linked_credit_note_nets_off() -> None:
    analysis = _analyze()
    assert ("within_tolerance", "d-in-12") in _codes(analysis)


def test_unlinked_credit_note_is_review() -> None:
    analysis = _analyze()
    items = [item for item in analysis.findings if item.code == "credit_note_unlinked"]
    assert len(items) == 1
    assert items[0].entity_id == "cn-us-01"
    assert items[0].entity == "credit_note"
    assert items[0].severity is Severity.REVIEW


# --- tax modes ------------------------------------------------------------------


def test_gross_with_tax_table() -> None:
    deal = _deal()
    deal["net_amount"] = "100000.00"
    deal["tax_name"] = "GST18"
    invoice = _invoice()
    invoice["net_amount"] = "118000.00"
    invoice["total"] = "118000.00"
    policy = MatchPolicy(tax="compare_gross_with_tax_table", tax_table={"GST18": Decimal("18.00")})
    analysis = analyze_recon(
        deals=[deal],
        invoices=[invoice],
        credit_notes=[],
        policy=policy,
        entities={"in-entity": EntityEntry(books_org_id_ref="in-entity", base_currency="INR")},
        org_ids={"in-entity": "555000001"},
    )
    assert {(item.code, item.entity_id) for item in analysis.findings} == {
        ("within_tolerance", "d-1")
    }


def test_unknown_tax_falls_back_to_net() -> None:
    deal = _deal()
    deal["tax_name"] = "MYSTERY"
    invoice = _invoice()
    policy = MatchPolicy(tax="compare_gross_with_tax_table", tax_table={"GST18": Decimal("18.00")})
    analysis = analyze_recon(
        deals=[deal],
        invoices=[invoice],
        credit_notes=[],
        policy=policy,
        entities={"in-entity": EntityEntry(books_org_id_ref="in-entity", base_currency="INR")},
        org_ids={"in-entity": "555000001"},
    )
    assert {(item.code, item.entity_id) for item in analysis.findings} == {
        ("within_tolerance", "d-1")
    }
    assert analysis.findings[0].evidence.get("tax_fallback_net") is True


# --- multi-org + cross-entity (UC-BK-3) -------------------------------------------


def test_org_fingerprint_is_stable_and_opaque() -> None:
    first = org_fingerprint("555000001")
    assert first == org_fingerprint("555000001")
    assert first.startswith("sha256:")
    assert "555000001" not in first
    assert org_fingerprint("555000002") != first


def test_cross_entity_invoice_carries_both_fingerprints() -> None:
    analysis = _analyze()
    items = [item for item in analysis.findings if item.code == "cross_entity_invoice"]
    by_entity = {(item.entity, item.entity_id) for item in items}
    assert ("deal", "d-in-09") in by_entity  # misfiled key match
    assert ("invoice", "i-us-02") in by_entity  # stray sharing the customer name
    deal_hit = next(item for item in items if item.entity_id == "d-in-09")
    assert deal_hit.evidence["entity_fingerprint"] == org_fingerprint("555000001")
    assert deal_hit.evidence["invoice_org_fingerprint"] == org_fingerprint("555000002")


def test_cross_entity_evidence_names_both_entities_without_empty_strings() -> None:
    analysis = _analyze()
    items = [item for item in analysis.findings if item.code == "cross_entity_invoice"]
    assert len(items) == 2
    for item in items:
        assert item.evidence.get("deal_entity"), "deal_entity must be present"
        assert item.evidence.get("invoice_entity"), "invoice_entity must be present"
        for value in item.evidence.values():
            if isinstance(value, str) and value.startswith("sha256:"):
                assert len(value) > len("sha256:")
            if isinstance(value, list):
                for entry in value:
                    assert entry != "", "no empty-string fingerprints"
        assert "" not in [value for value in item.evidence.values() if isinstance(value, str)]
        flat = [value for value in item.evidence.values() if isinstance(value, str)]
        flat += [
            entry for value in item.evidence.values() if isinstance(value, list) for entry in value
        ]
        assert "" not in flat
    deal_hit = next(item for item in items if item.entity_id == "d-in-09")
    assert deal_hit.evidence["deal_entity"] == "in-entity"
    assert deal_hit.evidence["invoice_entity"] == "us-entity"
    assert (
        deal_hit.message == "Invoice i-us-04 for deal d-in-09 (in-entity) was raised in us-entity."
    )
    stray = next(item for item in items if item.entity_id == "i-us-02")
    assert stray.evidence["invoice_entity"] == "us-entity"
    assert stray.evidence["deal_entity"] == "in-entity"
    assert "us-entity" in stray.message and "in-entity" in stray.message


def test_cross_entity_omits_fingerprints_without_profile_org_id() -> None:
    analysis = _analyze(org_ids={})
    items = [item for item in analysis.findings if item.code == "cross_entity_invoice"]
    assert items
    for item in items:
        assert item.evidence.get("deal_entity")
        assert item.evidence.get("invoice_entity")
        flat = [value for value in item.evidence.values() if isinstance(value, str)]
        assert "" not in flat
        assert "entity_fingerprint" not in item.evidence
        assert "invoice_org_fingerprint" not in item.evidence
    unavailable = next(item for item in analysis.findings if item.code == "source_unavailable")
    assert "org_fingerprint" not in unavailable.evidence


def test_wrong_entity_invoice_is_never_matched_by_name() -> None:
    analysis = _analyze()
    assert ("within_tolerance", "d-in-09") not in _codes(analysis)
    assert ("missing_invoice", "d-in-09") in _codes(analysis)


def test_report_carries_no_raw_org_id() -> None:
    from zohokit.reports import render_html
    from zohokit.reports.xlsx import render_xlsx

    ctx = RunContext(now=datetime(2026, 10, 1, tzinfo=UTC), mode="offline")
    from zohokit.modules.books.reconcile import run_recon

    data = _marigold_data()
    report = run_recon(
        data,
        policy=_policy(),
        entities=_entities(),
        org_ids=ORG_IDS,
        fx_rates=load_fx_rates(MARIGOLD / "fx_rates.csv"),
        unavailable=("eu-entity",),
        ctx=ctx,
    )
    blob = render_json(report) + render_html(report)
    for raw in ORG_IDS.values():
        assert raw not in blob
    assert org_fingerprint("555000001") in blob
    workbook = render_xlsx(report)
    for raw in ORG_IDS.values():
        assert raw.encode("utf-8") not in workbook
    from openpyxl import load_workbook

    book = load_workbook(filename=__import__("io").BytesIO(workbook), data_only=False)
    cells = [
        str(cell.value)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    for raw in ORG_IDS.values():
        assert not any(raw in text for text in cells)
    assert any(org_fingerprint("555000001") in text for text in cells)


def test_crm_criteria_matching() -> None:
    entry = EntityEntry(
        books_org_id_ref="in-entity",
        base_currency="INR",
        crm_criteria={"field": "Entity", "equals": "India"},
    )
    assert deal_in_entity({"Entity": "India"}, entry) is True
    assert deal_in_entity({"Entity": "US"}, entry) is False
    assert deal_in_entity({}, EntityEntry(books_org_id_ref="x", base_currency="INR")) is True


# --- FX --------------------------------------------------------------------------


def test_missing_rate_is_review_never_guessed() -> None:
    analysis = _analyze()
    items = [item for item in analysis.findings if item.code == "fx_rate_missing"]
    assert [(item.entity_id, item.severity) for item in items] == [("d-in-10", Severity.REVIEW)]
    assert ("within_tolerance", "d-in-10") not in _codes(analysis)
    assert ("outside_tolerance", "d-in-10") not in _codes(analysis)


def test_converted_amount_shown_next_to_original() -> None:
    analysis = _analyze()
    item = next(item for item in analysis.findings if item.entity_id == "d-in-11")
    assert item.code == "within_tolerance"
    (conversion,) = item.evidence["conversions"]
    assert conversion["original"] == "1000.00"
    assert conversion["original_currency"] == "USD"
    assert conversion["converted"] == "88410.00"
    assert conversion["converted_currency"] == "INR"
    assert conversion["rate"] == "88.4100"


def test_money_uses_decimal_not_float() -> None:
    from zohokit.modules.books.fx import FxRate

    rate = FxRate(
        on=__import__("datetime").date(2026, 9, 20),
        frm="USD",
        to="INR",
        rate=Decimal("2.675"),
        source="test",
    )
    assert convert(Decimal("1.00"), rate) == Decimal("2.675")
    from zohokit.core.money import quantize_money

    assert quantize_money(convert(Decimal("1.00"), rate)) == Decimal("2.68")
    # The float path rounds the same value down: binary float cannot hold
    # 2.675, so round() sees 2.67499... and returns 2.67.
    assert round(1.00 * 2.675, 2) == 2.67


def test_fx_lookup_is_exact_only() -> None:
    from datetime import date as _date

    rates = load_fx_rates(MARIGOLD / "fx_rates.csv")
    assert lookup_rate(rates, on=_date(2026, 9, 20), frm="USD", to="INR") is not None
    assert lookup_rate(rates, on=_date(2026, 9, 21), frm="USD", to="INR") is None
    assert lookup_rate(rates, on=_date(2026, 9, 20), frm="INR", to="USD") is None


# --- statuses ---------------------------------------------------------------------


def test_draft_excluded_by_default_with_info() -> None:
    analysis = _analyze()
    assert ("draft_invoice_excluded", "d-in-07") in _codes(analysis)
    assert ("missing_invoice", "d-in-07") in _codes(analysis)


def test_void_excluded_by_default() -> None:
    analysis = _analyze()
    codes = {(item.code, item.entity_id) for item in analysis.findings}
    assert ("void_invoice_excluded", "d-in-16") in codes
    assert ("missing_invoice", "d-in-16") in codes
    void_hit = next(
        item
        for item in analysis.findings
        if item.code == "void_invoice_excluded" and item.entity_id == "d-in-16"
    )
    assert void_hit.severity is Severity.INFO
    assert void_hit.evidence["invoice_id"] == "i-in-16"
    assert void_hit.message == "Invoice i-in-16 for deal d-in-16 is void, excluded by policy."
    assert not [item for item in analysis.findings if item.entity_id == "i-in-16"]


def test_void_excluded_silently_before_fix_now_audited() -> None:
    data = _marigold_data()
    invoice = dict(data["invoices"][0])
    invoice["id"] = "i-void-1"
    invoice["status"] = "void"
    analysis = _analyze(invoices=[invoice])
    codes = {(item.code, item.entity_id) for item in analysis.findings}
    assert ("missing_invoice", "d-in-01") in codes
    assert ("void_invoice_excluded", "d-in-01") in codes


# --- period windowing ---------------------------------------------------------------


def test_record_date_uses_own_timezone() -> None:
    from datetime import date as _date

    assert record_date("2026-09-30T23:30:00+05:30") == _date(2026, 9, 30)
    assert record_date("2026-10-01T00:30:00+05:30") == _date(2026, 10, 1)
    assert record_date("2026-08-31T19:00:00-07:00") == _date(2026, 8, 31)
    assert record_date("2026-09-01") == _date(2026, 9, 1)
    assert record_date("") is None


def test_excluded_boundaries_have_no_findings() -> None:
    from datetime import date as _date

    assert in_window("2026-09-30T23:30:00+05:30", start=_date(2026, 9, 1), end=_date(2026, 9, 30))
    assert not in_window(
        "2026-10-01T00:30:00+05:30", start=_date(2026, 9, 1), end=_date(2026, 9, 30)
    )
    assert not in_window(
        "2026-08-31T19:00:00-07:00", start=_date(2026, 9, 1), end=_date(2026, 9, 30)
    )
    analysis = _analyze()
    codes = _codes(analysis)
    assert not [(code, entity) for code, entity in codes if entity in ("d-in-13", "d-in-14")]
    assert ("within_tolerance", "d-in-15") in codes


def test_invoice_date_mode_windows_invoices() -> None:
    from datetime import date as _date

    from zohokit.modules.books.policy import PeriodWindow

    policy = MatchPolicy(
        period=PeriodWindow(by="invoice_date", start=_date(2026, 9, 1), end=_date(2026, 9, 30))
    )
    analysis = _analyze(policy=policy, unavailable=())
    codes = {(item.code, item.entity_id) for item in analysis.findings}
    assert ("within_tolerance", "d-in-01") in codes


# --- robustness ---------------------------------------------------------------------


def test_invalid_amount_isolated_like_legacy() -> None:
    analysis = _analyze()
    assert ("invalid_amount", "d-in-08") in _codes(analysis)
    assert ("within_tolerance", "d-in-01") in _codes(analysis)


def test_invalid_date_is_review_and_excludes_row() -> None:
    analysis = _analyze()
    codes = _codes(analysis)
    assert ("invalid_date", "d-in-17") in codes
    assert ("invalid_date", "i-in-99") in codes
    # The bad-date deal is excluded from matching (no amount verdict, no
    # missing_invoice); the bad-date invoice never matches its deal.
    assert ("missing_invoice", "d-in-17") not in codes
    assert ("within_tolerance", "d-in-17") not in codes
    assert ("missing_invoice", "d-in-06") in codes
    assert ("within_tolerance", "d-in-01") in codes
    deal_hit = next(
        item
        for item in analysis.findings
        if item.code == "invalid_date" and item.entity_id == "d-in-17"
    )
    assert deal_hit.severity is Severity.REVIEW
    assert deal_hit.entity == "deal"
    invoice_hit = next(
        item
        for item in analysis.findings
        if item.code == "invalid_date" and item.entity_id == "i-in-99"
    )
    assert invoice_hit.entity == "invoice"
    assert invoice_hit.severity is Severity.REVIEW


def test_malformed_org_gives_source_unavailable_while_others_continue() -> None:
    analysis = _analyze()
    codes = _codes(analysis)
    assert ("source_unavailable", "eu-entity") in codes
    assert ("within_tolerance", "d-in-01") in codes
    assert not [(code, entity) for code, entity in codes if entity == "d-eu-01"]
    item = next(item for item in analysis.findings if item.code == "source_unavailable")
    assert item.entity == "entity"
    assert item.severity is Severity.ERROR
    assert item.evidence["org_fingerprint"] == org_fingerprint("555000003")


# --- identity stability ---------------------------------------------------------------


def test_marigold_ids_unique() -> None:
    analysis = _analyze()
    ids = [item.id for item in analysis.findings]
    assert len(ids) == len(set(ids))


def _ids_for(deals: object, invoices: object) -> set[str]:
    analysis = _analyze(deals=deals, invoices=invoices)  # type: ignore[arg-type]
    return {item.id for item in analysis.findings}


@given(st.data())
def test_ids_stable_under_shuffle(data: st.DataObject) -> None:
    payload = _marigold_data()
    deals = [
        payload["deals"][index]
        for index in data.draw(st.permutations(range(len(payload["deals"]))))
    ]
    invoices = [
        payload["invoices"][index]
        for index in data.draw(st.permutations(range(len(payload["invoices"]))))
    ]
    assert _ids_for(payload["deals"], payload["invoices"]) == _ids_for(deals, invoices)


def test_ids_stable_when_rows_prepended() -> None:
    payload = _marigold_data()
    before = _ids_for(payload["deals"], payload["invoices"])
    extra_deal = {
        "id": "d-extra",
        "entity": "in-entity",
        "currency": "INR",
        "net_amount": "10.00",
        "customer_name": "Extra Co",
        "closing_date": "2026-09-05",
        "reference_number": "REF-X",
    }
    extra_invoice = {
        "id": "i-extra",
        "entity": "in-entity",
        "currency": "INR",
        "net_amount": "10.00",
        "customer_name": "Extra Co",
        "status": "sent",
        "date": "2026-09-05",
        "reference_number": "INV-X",
        "custom_fields": [{"label": "CRM Deal ID", "value": "d-extra"}],
    }
    after = _ids_for([extra_deal, *payload["deals"]], [extra_invoice, *payload["invoices"]])
    assert before <= after


# --- CLI (v2 flags) -------------------------------------------------------------------


def test_cli_v2_offline_recon(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    out = tmp_path / "report.json"
    result = CliRunner().invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "legacy-v1",
            "--policy",
            str(MARIGOLD / "policy.yaml"),
            "--entity-map",
            str(MARIGOLD / "entity_map.yaml"),
            "--fx-rates",
            str(MARIGOLD / "fx_rates.csv"),
            "--unavailable",
            "eu-entity",
            "--out",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["module"] == "books"
    assert {item["code"] for item in payload["findings"]} >= {
        "under_invoiced",
        "over_invoiced",
        "credit_note_unlinked",
        "fx_rate_missing",
        "draft_invoice_excluded",
        "source_unavailable",
        "outside_tolerance",
        "within_tolerance",
        "invalid_amount",
    }


def test_cli_v2_auto_detects_recon_envelope(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    out = tmp_path / "report.json"
    result = CliRunner().invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--policy",
            str(MARIGOLD / "policy.yaml"),
            "--entity-map",
            str(MARIGOLD / "entity_map.yaml"),
            "--fx-rates",
            str(MARIGOLD / "fx_rates.csv"),
            "--unavailable",
            "eu-entity",
            "--out",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["module"] == "books"
    assert {item["code"] for item in payload["findings"]} >= {
        "void_invoice_excluded",
        "invalid_date",
    }


def test_cli_v2_accepts_recon_v2_format(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    out = tmp_path / "report.json"
    result = CliRunner().invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "recon-v2",
            "--policy",
            str(MARIGOLD / "policy.yaml"),
            "--entity-map",
            str(MARIGOLD / "entity_map.yaml"),
            "--out",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["module"] == "books"


def test_cli_recon_v2_without_policy_exits_1(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    out = tmp_path / "report.json"
    result = CliRunner().invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "recon-v2",
            "--out",
            str(out),
        ],
    )
    assert result.exit_code == 1


def test_cli_bad_period_and_missing_map_exit_1() -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    runner = CliRunner()
    bad_period = runner.invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "legacy-v1",
            "--policy",
            str(MARIGOLD / "policy.yaml"),
            "--entity-map",
            str(MARIGOLD / "entity_map.yaml"),
            "--period",
            "2026-13",
        ],
    )
    assert bad_period.exit_code == 1
    no_map = runner.invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "legacy-v1",
            "--period",
            "2026-09",
        ],
    )
    assert no_map.exit_code == 1


def test_cli_live_needs_experimental() -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    result = CliRunner().invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "legacy-v1",
            "--live",
            "--profile",
            "dev-in",
            "--books-orgs",
            "in-entity",
            "--entity-map",
            str(MARIGOLD / "entity_map.yaml"),
        ],
    )
    assert result.exit_code == 1
    assert "experimental" in result.output.lower()


# --- Marigold answer key ---------------------------------------------------------------


def test_finance_readable_messages_exact() -> None:
    analysis = _analyze()
    messages = {(item.code, item.entity_id): item.message for item in analysis.findings}
    assert messages[("outside_tolerance", "d-in-03")] == (
        "Invoiced 49,500.00 INR against a deal net of 50,000.00 INR "
        "(-500.00), outside tolerance (\u00b11.00 and 0.5%)."
    )
    assert messages[("under_invoiced", "d-in-04")] == (
        "Deal d-in-04 under-invoiced: staged invoices total 140,000.00 INR "
        "against a deal net of 200,000.00 INR (-60,000.00), outside tolerance "
        "(\u00b11.00 and 0.5%)."
    )
    assert messages[("over_invoiced", "d-in-05")] == (
        "Deal d-in-05 over-invoiced: staged invoices total 120,000.00 INR "
        "against a deal net of 100,000.00 INR (+20,000.00), outside tolerance "
        "(\u00b11.00 and 0.5%)."
    )
    assert messages[("fx_rate_missing", "d-in-10")] == (
        "Deal d-in-10: no exact USD->INR rate for 2026-09-15; comparison skipped, never guessed."
    )
    assert messages[("credit_note_unlinked", "cn-us-01")] == (
        "Credit note cn-us-01 (400.00 USD) links to no known invoice."
    )
    assert messages[("draft_invoice_excluded", "d-in-07")] == (
        "Invoice i-in-07 for deal d-in-07 is a draft, excluded by policy."
    )
    assert messages[("void_invoice_excluded", "d-in-16")] == (
        "Invoice i-in-16 for deal d-in-16 is void, excluded by policy."
    )
    assert messages[("within_tolerance", "d-in-01")] == (
        "Deal d-in-01 matched: invoiced 420,000.00 INR against a deal net of "
        "420,000.00 INR, within tolerance (\u00b11.00 and 0.5%)."
    )
    assert messages[("missing_invoice", "d-in-06")] == (
        "Deal d-in-06 (75,000.00 INR) is missing its invoice for the period."
    )
    assert messages[("orphan_invoice_reference", "i-us-03")] == (
        "Invoice i-us-03 (2,500.00 USD) references no known deal."
    )
    assert messages[("invalid_amount", "d-in-08")] == (
        "Row d-in-08 carries a malformed amount '12.34.56'; excluded from amount comparison."
    )
    assert messages[("invalid_date", "d-in-17")] == (
        "Row d-in-17 carries a missing or unparseable closing_date ''; excluded from matching."
    )
    assert messages[("source_unavailable", "eu-entity")] == (
        "Org pull for eu-entity failed; its records are out of scope while others continue."
    )


def test_marigold_answer_key_exact_set() -> None:
    analysis = _analyze()
    assert _codes(analysis) == {
        ("credit_note_unlinked", "cn-us-01"),
        ("cross_entity_invoice", "d-in-09"),
        ("cross_entity_invoice", "i-us-02"),
        ("draft_invoice_excluded", "d-in-07"),
        ("fx_rate_missing", "d-in-10"),
        ("invalid_amount", "d-in-08"),
        ("invalid_date", "d-in-17"),
        ("invalid_date", "i-in-99"),
        ("missing_invoice", "d-in-06"),
        ("missing_invoice", "d-in-07"),
        ("missing_invoice", "d-in-08"),
        ("missing_invoice", "d-in-09"),
        ("missing_invoice", "d-in-16"),
        ("missing_invoice", "d-us-02"),
        ("orphan_invoice_reference", "i-us-03"),
        ("outside_tolerance", "d-in-03"),
        ("over_invoiced", "d-in-05"),
        ("source_unavailable", "eu-entity"),
        ("under_invoiced", "d-in-04"),
        ("void_invoice_excluded", "d-in-16"),
        ("within_tolerance", "d-in-01"),
        ("within_tolerance", "d-in-02"),
        ("within_tolerance", "d-in-11"),
        ("within_tolerance", "d-in-12"),
        ("within_tolerance", "d-in-15"),
        ("within_tolerance", "d-us-01"),
    }
