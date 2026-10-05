"""Reconcile workbook tests: structure golden + findings (TK-MIG-F10)."""

from __future__ import annotations

import io
import json
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from zohokit.cli import app
from zohokit.modules.migration.mapping import MappingDoc
from zohokit.modules.migration.reconcile import (
    RECONCILE_SHEETS,
    recap_findings,
    reconcile_entity,
    render_reconcile_workbook,
)

ROOT = Path(__file__).resolve().parent.parent.parent.parent
FIXTURES = ROOT / "fixtures" / "migration"

runner = CliRunner()


def _doc():  # type: ignore[no-untyped-def]
    return MappingDoc.model_validate(
        {
            "version": 1,
            "source": "pipedrive",
            "entities": [
                {
                    "name": "deals",
                    "source_kind": "deals",
                    "target_module": "Deals",
                    "fields": {
                        "Deal_Name": {"from": "Title"},
                        "Stage": {"from": "Stage"},
                    },
                    "external_id": {"field": "External_ID__s", "from": "ID"},
                    "lookups": {"Account_Name": {"entity": "companies", "via": "Org"}},
                },
                {
                    "name": "companies",
                    "source_kind": "organizations",
                    "target_module": "Accounts",
                    "fields": {"Account_Name": {"from": "Name"}},
                    "external_id": {"field": "External_ID__s", "from": "ID"},
                },
            ],
        }
    )


def _rows():  # type: ignore[no-untyped-def]
    source = [
        {"ID": "1", "Title": "Pilot", "Stage": "Qualification", "Org": "10"},
        {"ID": "2", "Title": "Renewal", "Stage": "Closed Won", "Org": "11"},
        {"ID": "3", "Title": "Ghost", "Stage": "Qualification", "Org": "99"},
    ]
    target = [
        {"External_ID__s": "1", "Deal_Name": "Pilot", "Stage": "Qualification"},
        {"External_ID__s": "2", "Deal_Name": "Renewal CHANGED", "Stage": "Closed Won"},
        {"External_ID__s": "9", "Deal_Name": "Spare", "Stage": "Qualification"},
    ]
    parents = {"companies": {"10", "11"}}
    return source, target, parents


def test_reconcile_counts_diffs_and_gaps() -> None:
    doc = _doc()
    source, target, parents = _rows()
    recap = reconcile_entity(
        doc.entities[0], source, ["External_ID__s", "Deal_Name", "Stage"], target, parents, seed=42
    )
    assert (recap.source_total, recap.target_total) == (3, 3)
    assert recap.missing_keys == ["3"]
    assert recap.extra_keys == ["9"]
    assert recap.sample_diffs == [("2", "Deal_Name", "Renewal", "Renewal CHANGED")]
    assert recap.relationship_gaps == [("3", "Org", "99")]
    assert ("Qualification", 2, 2) in recap.stage_rows
    findings = recap_findings([recap])
    assert sorted(item.code for item in findings) == [
        "reconcile_count_mismatch",
        "reconcile_field_diff",
        "reconcile_relationship_gap",
    ]
    assert [item.severity for item in findings] == ["error", "review", "error"]


def test_workbook_structure_golden() -> None:
    doc = _doc()
    source, target, parents = _rows()
    recap = reconcile_entity(
        doc.entities[0], source, ["External_ID__s", "Deal_Name", "Stage"], target, parents, seed=7
    )
    book = load_workbook(filename=io.BytesIO(render_reconcile_workbook([recap], seed=7)))
    assert book.sheetnames == list(RECONCILE_SHEETS)
    summary = book["Summary"]
    assert summary["A1"].value == "Migration reconciliation"
    assert summary["B2"].value == 7
    assert summary.freeze_panes == "A5"
    assert [summary.cell(row=5, column=c).value for c in range(1, 7)] == [
        "deals",
        "Deals",
        3,
        3,
        1,
        1,
    ]
    counts = book["Counts"]
    assert counts["A1"].value == "Entity"
    assert counts.freeze_panes == "A2"
    delta = [counts.cell(row=r, column=5).value for r in range(2, 4)]
    assert delta == ["=C2-D2", "=C3-D3"]
    diffs = book["Sample diffs"]
    assert [diffs.cell(row=1, column=c).value for c in range(1, 7)] == [
        "Entity",
        "Key",
        "Field",
        "Source",
        "Target",
        "Match",
    ]
    assert diffs.cell(row=2, column=2).value == "2"
    assert diffs.freeze_panes == "A2"
    relations = book["Relationships"]
    assert relations["A1"].value == "Entity"
    assert relations.cell(row=2, column=2).value == "3"
    assert relations.freeze_panes == "A2"


def test_cli_reconcile_writes_workbook(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    (source / "persons.csv").write_text("ID,Name\n1,Aarav\n2,Meera\n", encoding="utf-8")
    target = tmp_path / "target"
    target.mkdir()
    (target / "Contacts.csv").write_text("External_ID__s,Last_Name\n1,Aarav\n", encoding="utf-8")
    mapping = tmp_path / "mapping.yaml"
    mapping.write_text(
        "version: 1\n"
        "source: generic\n"
        "entities:\n"
        "  - name: people\n"
        "    source_kind: persons\n"
        "    target_module: Contacts\n"
        "    fields:\n"
        "      Last_Name: {from: Name}\n"
        "    external_id: {field: External_ID__s, from: ID}\n",
        encoding="utf-8",
    )
    workbook = tmp_path / "reconcile.xlsx"
    report_path = tmp_path / "report.json"
    result = runner.invoke(
        app,
        [
            "migration",
            "reconcile",
            "--mapping",
            str(mapping),
            "--source",
            str(source),
            "--target",
            str(target),
            "--workbook",
            str(workbook),
            "--format",
            "json",
            "--out",
            str(report_path),
        ],
    )
    assert result.exit_code == 0, result.output
    book = load_workbook(filename=io.BytesIO(workbook.read_bytes()))
    assert book.sheetnames == list(RECONCILE_SHEETS)
    payload = json.loads(report_path.read_text(encoding="utf-8"))
    assert payload["ready"] is False
    assert [item["code"] for item in payload["findings"]] == ["reconcile_count_mismatch"]
