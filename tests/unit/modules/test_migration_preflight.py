"""Preflight + CLI tests (TK-MIG-F1..F3, TK-CORE-8 xlsx)."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from zohokit.cli import app
from zohokit.core.context import RunContext
from zohokit.modules.migration.mapping import load_mapping
from zohokit.modules.migration.metadata import metadata_from_dir
from zohokit.modules.migration.preflight import SourceSpec, run_preflight

ROOT = Path(__file__).resolve().parent.parent.parent.parent
FIXTURES = ROOT / "fixtures" / "migration"

runner = CliRunner()


def _offline_report():  # type: ignore[no-untyped-def]
    doc = load_mapping(FIXTURES / "mapping.yaml")
    sources = {
        entity.name: SourceSpec(
            path=str(FIXTURES / "source" / f"{entity.source_kind}.csv"),
            kind=entity.source_kind,
        )
        for entity in doc.entities
    }
    metadata = metadata_from_dir(FIXTURES / "fields")
    return run_preflight(
        doc, sources, metadata, ctx=RunContext(now=datetime(2026, 1, 5, tzinfo=UTC))
    )


def test_clean_preflight_is_ready() -> None:
    report = _offline_report()
    assert report.ready is True
    assert report.summary.error == 0
    assert report.findings == []


def test_row_parse_error_never_aborts(tmp_path: Path) -> None:
    doc = load_mapping(FIXTURES / "mapping.yaml")
    bad_persons = tmp_path / "persons.csv"
    bad_persons.write_text(
        "ID,Name,Email,Phone,Organization ID,Source,Created\n"
        "1,Aarav Sharma,aarav@example.invalid,+919000000001,10,web,2026-01-05\n"
        "2,Short row\n"
        "\n"
        "3,Zara Khan,zara@example.invalid,+919000000005,10,web,2026-01-06\n",
        encoding="utf-8",
    )
    sources = {
        entity.name: SourceSpec(
            path=str(FIXTURES / "source" / f"{entity.source_kind}.csv"),
            kind=entity.source_kind,
        )
        for entity in doc.entities
    }
    sources["people"] = SourceSpec(path=str(bad_persons), kind="persons")
    metadata = metadata_from_dir(FIXTURES / "fields")
    report = run_preflight(
        doc, sources, metadata, ctx=RunContext(now=datetime(2026, 1, 5, tzinfo=UTC))
    )
    codes = sorted(item.code for item in report.findings)
    assert codes == ["row_parse_error", "row_parse_error"]
    reasons = sorted(item.evidence["reason"] for item in report.findings)
    assert reasons == ["blank_row", "wrong_column_count"]
    assert report.ready is False


def test_cli_preflight_json_and_xlsx(tmp_path: Path) -> None:
    result = runner.invoke(
        app,
        [
            "migration",
            "preflight",
            "--mapping",
            str(FIXTURES / "mapping.yaml"),
            "--source",
            str(FIXTURES / "source"),
            "--fields-dir",
            str(FIXTURES / "fields"),
            "--format",
            "json",
        ],
    )
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["ready"] is True
    out = tmp_path / "report.xlsx"
    result = runner.invoke(
        app,
        [
            "migration",
            "preflight",
            "--mapping",
            str(FIXTURES / "mapping.yaml"),
            "--source",
            str(FIXTURES / "source"),
            "--fields-dir",
            str(FIXTURES / "fields"),
            "--format",
            "xlsx",
            "--out",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    book = load_workbook(filename=io.BytesIO(out.read_bytes()))
    assert book.sheetnames == ["Summary", "Findings", "Sign-off"]


def test_cli_bad_mapping_exits_1_with_line(tmp_path: Path) -> None:
    mapping = tmp_path / "mapping.yaml"
    mapping.write_text(
        "version: 1\n"
        "entities:\n"
        "  - name: people\n"
        "    source_kind: persons\n"
        "    target_module: Contacts\n"
        "    fields:\n"
        "      Last_Name: {from: Name, transform: [bogus]}\n",
        encoding="utf-8",
    )
    result = runner.invoke(
        app,
        [
            "migration",
            "preflight",
            "--mapping",
            str(mapping),
            "--source",
            str(FIXTURES / "source"),
            "--fields-dir",
            str(FIXTURES / "fields"),
        ],
    )
    assert result.exit_code == 1
    assert ":7: " in result.output


def test_cli_needs_metadata_or_live() -> None:
    result = runner.invoke(
        app,
        [
            "migration",
            "preflight",
            "--mapping",
            str(FIXTURES / "mapping.yaml"),
            "--source",
            str(FIXTURES / "source"),
        ],
    )
    assert result.exit_code == 1
    assert "--fields-dir" in result.output


def test_cli_legacy_path_unchanged() -> None:
    fixture = ROOT / "tests" / "golden" / "legacy" / "migration" / "inputs" / "examples.json"
    result = runner.invoke(app, ["migration", "audit", str(fixture), "--format", "json"])
    assert result.exit_code == 0
    assert json.loads(result.output)["ready"] is False
