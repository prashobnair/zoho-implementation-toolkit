"""Finance workbook tests (TK-BK-F7) + formula-injection guard (STD §4 safety)."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from zohokit.core.context import RunContext
from zohokit.modules.books.entities import load_entity_map
from zohokit.modules.books.fx import load_fx_rates
from zohokit.modules.books.policy import load_policy
from zohokit.modules.books.reconcile import analyze_recon, run_recon
from zohokit.modules.books.workbook import (
    BOOKS_SHEETS,
    render_books_workbook,
    sum_range,
)
from zohokit.reports.xlsx import render_xlsx, safe_text

ROOT = Path(__file__).resolve().parent.parent.parent.parent
MARIGOLD = ROOT / "fixtures" / "books" / "marigold"
ORG_IDS = {"in-entity": "555000001", "us-entity": "555000002", "eu-entity": "555000003"}

PAYLOAD = '=HYPERLINK("https://x.invalid","click")'


def _marigold_report() -> object:
    from zohokit.core.findings import Report as _Report

    data = json.loads((MARIGOLD / "recon.json").read_text(encoding="utf-8"))
    report = run_recon(
        data,
        policy=load_policy(MARIGOLD / "policy.yaml"),
        entities=load_entity_map(MARIGOLD / "entity_map.yaml"),
        org_ids=ORG_IDS,
        fx_rates=load_fx_rates(MARIGOLD / "fx_rates.csv"),
        unavailable=("eu-entity",),
        ctx=RunContext(now=datetime(2026, 10, 1, tzinfo=UTC), mode="offline"),
    )
    assert isinstance(report, _Report)
    return report


def _book(data: bytes):  # type: ignore[no-untyped-def]
    return load_workbook(filename=io.BytesIO(data), data_only=False)


def test_workbook_tabs_in_order() -> None:
    book = _book(render_books_workbook(_marigold_report()))
    assert (
        book.sheetnames
        == list(BOOKS_SHEETS)
        == [
            "Matched",
            "Mismatched",
            "Missing invoice",
            "Orphan invoice",
            "Needs review",
            "FX",
            "Sign-off",
        ]
    )


def test_totals_are_formulas_that_recompute() -> None:
    book = _book(render_books_workbook(_marigold_report()))
    matched = book["Matched"]
    totals = [cell.value for cell in matched[matched.max_row] if cell.value is not None]
    formulas = [value for value in totals if isinstance(value, str) and value.startswith("=")]
    assert formulas, "totals row must carry SUM formulas"
    for formula in formulas:
        assert formula.startswith("=SUM(")
    net_total = sum_range(matched.cell(row=matched.max_row, column=5).value, matched)
    inv_total = sum_range(matched.cell(row=matched.max_row, column=6).value, matched)
    assert net_total == pytest.approx(775410.00)
    assert inv_total == pytest.approx(775410.00)
    mismatched = book["Mismatched"]
    mis_net = sum_range(mismatched.cell(row=mismatched.max_row, column=5).value, mismatched)
    assert mis_net == pytest.approx(50000.00 + 200000.00 + 100000.00)
    fx_sheet = book["FX"]
    fx_total = sum_range(fx_sheet.cell(row=fx_sheet.max_row, column=5).value, fx_sheet)
    assert fx_total == pytest.approx(88410.00)


def test_headers_frozen_and_signoff_blank() -> None:
    book = _book(render_books_workbook(_marigold_report()))
    for name in BOOKS_SHEETS[:6]:
        assert book[name].freeze_panes == "A2"
    signoff = book["Sign-off"]
    assert signoff["B3"].value in (None, "")
    assert signoff["B4"].value in (None, "")
    assert signoff["B5"].value in (None, "")


def test_safe_text_guards_formula_prefixes() -> None:
    assert safe_text(PAYLOAD) == "'" + PAYLOAD
    assert safe_text("+cmd").startswith("'")
    assert safe_text("-cmd").startswith("'")
    assert safe_text("@cmd").startswith("'")
    assert safe_text("\tindented").startswith("'")
    assert safe_text("plain") == "plain"
    assert safe_text(420000.0) == 420000.0
    assert safe_text("") == ""


def _payload_report() -> object:
    data = json.loads((MARIGOLD / "recon.json").read_text(encoding="utf-8"))
    data["deals"] = [dict(data["deals"][0])]
    data["deals"][0]["id"] = PAYLOAD
    data["deals"][0]["customer_name"] = PAYLOAD
    data["invoices"] = [dict(data["invoices"][0])]
    data["invoices"][0]["custom_fields"] = [{"label": "CRM Deal ID", "value": PAYLOAD}]
    data["invoices"][0]["customer_name"] = PAYLOAD
    data["credit_notes"] = []
    analysis = analyze_recon(
        deals=data["deals"],
        invoices=data["invoices"],
        credit_notes=[],
        policy=load_policy(MARIGOLD / "policy.yaml"),
        entities={"in-entity": load_entity_map(MARIGOLD / "entity_map.yaml")["in-entity"]},
        org_ids={"in-entity": "555000001"},
    )
    from zohokit.modules.books.report import build_report

    return build_report(
        analysis,
        ctx=RunContext(now=datetime(2026, 10, 1, tzinfo=UTC), mode="offline"),
        inputs_sha256="payload-case",
    )


def test_injection_stored_as_text_in_every_workbook() -> None:
    report = _payload_report()
    for blob in (render_xlsx(report), render_books_workbook(report)):
        book = _book(blob)
        formula_cells = []
        guarded = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_cells.append((sheet.title, cell.coordinate, cell.value))
                    elif isinstance(cell.value, str) and cell.value.startswith("'="):
                        guarded.append(cell.value)
        assert guarded, "payload must be stored as guarded text"
        assert all(value == "'" + PAYLOAD for value in guarded)
        for _title, _coord, value in formula_cells:
            assert isinstance(value, str) and value.startswith("=SUM("), value


def test_cli_xlsx_writes_finance_workbook(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from zohokit.cli import app

    out = tmp_path / "recon.xlsx"
    result = CliRunner().invoke(
        app,
        [
            "books",
            "reconcile",
            str(MARIGOLD / "recon.json"),
            "--input-format",
            "legacy-v1",
            "--policy",
            str(MARIGOLD / "policy.yaml"),
            "--entity-map",
            str(MARIGOLD / "entity_map.yaml"),
            "--fx-rates",
            str(MARIGOLD / "fx_rates.csv"),
            "--unavailable",
            "eu-entity",
            "--format",
            "xlsx",
            "--out",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    book = _book(out.read_bytes())
    assert book.sheetnames == list(BOOKS_SHEETS)
