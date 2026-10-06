"""Post-import reconciliation workbook (TK-MIG-F10).

Compares the source export against target dumps (``--target`` holding
``<TargetModule>.csv`` files): count parity per module, a seeded sample
of field-level compares with mapping transforms applied, relationship
integrity through the mapped lookups, and stage distribution deltas.
Output is an XLSX workbook (Summary / Counts / Sample diffs /
Relationships) plus standard findings:

- ``reconcile_count_mismatch`` (error) per module with missing/extra rows,
- ``reconcile_field_diff`` (review) per sampled row with a field diff,
- ``reconcile_relationship_gap`` (error) per orphaned child row.

Finding evidence stays value-free (field names and counts only); the
workbook's Sample diffs tab carries the values for the human reviewer.
"""

from __future__ import annotations

import io
import random
from dataclasses import dataclass
from dataclasses import field as dataclass_field

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from zohokit.connectors.sources.csv_reader import open_csv
from zohokit.core.findings import Finding, Severity
from zohokit.modules.migration import dedupe as dedupe_mod
from zohokit.modules.migration.mapping import EntityMapping, MappingDoc
from zohokit.modules.migration.metadata import mapped_value

#: Sampled rows per entity (seeded; deterministic for one seed).
RECONCILE_SAMPLE = 50

#: Workbook tabs in order (asserted by the golden structure test).
RECONCILE_SHEETS = ("Summary", "Counts", "Sample diffs", "Relationships")

_BOLD = Font(bold=True)


@dataclass
class EntityRecap:
    """Reconciliation outcome for one entity."""

    entity: str
    module: str
    source_total: int = 0
    target_total: int = 0
    missing_keys: list[str] = dataclass_field(default_factory=list)
    extra_keys: list[str] = dataclass_field(default_factory=list)
    sample_diffs: list[tuple[str, str, str, str]] = dataclass_field(default_factory=list)
    relationship_gaps: list[tuple[str, str, str]] = dataclass_field(default_factory=list)
    stage_rows: list[tuple[str, int, int]] = dataclass_field(default_factory=list)
    notes: list[str] = dataclass_field(default_factory=list)


def _read_csv_rows(path: str) -> tuple[list[str], list[dict[str, str]]]:
    """Header plus all data rows (target dumps are small by construction)."""
    stream = open_csv(path)
    try:
        header = list(stream.dialect.header)
        rows = [row for _, row, issue in stream.rows() if issue is None and row is not None]
    finally:
        stream.close()
    return header, rows


def _external_key(entity: EntityMapping, row: dict[str, str]) -> str | None:
    """Join key: the external-ID source value, if the mapping declares one."""
    if entity.external_id is None:
        return None
    value = (row.get(entity.external_id.from_col) or "").strip()
    return value or None


def reconcile_entity(
    entity: EntityMapping,
    source_rows: list[dict[str, str]],
    target_header: list[str],
    target_rows: list[dict[str, str]],
    parent_keys: dict[str, set[str]],
    *,
    seed: int,
) -> EntityRecap:
    """Reconcile one entity's source rows against its target dump."""
    recap = EntityRecap(entity=entity.name, module=entity.target_module)
    recap.source_total = len(source_rows)
    recap.target_total = len(target_rows)
    if entity.external_id is None:
        recap.notes.append("no external_id mapping: sample join skipped")
        return recap
    external_field = entity.external_id.field
    if external_field not in target_header:
        recap.notes.append(f"target lacks external ID column {external_field}: join skipped")
        return recap
    source_by_key: dict[str, dict[str, str]] = {}
    for row in source_rows:
        key = _external_key(entity, row)
        if key is not None:
            source_by_key[key] = row
    target_by_key: dict[str, dict[str, str]] = {}
    for row in target_rows:
        key = (row.get(external_field) or "").strip()
        if key:
            target_by_key[key] = row
    recap.missing_keys = sorted(set(source_by_key) - set(target_by_key))
    recap.extra_keys = sorted(set(target_by_key) - set(source_by_key))
    comparable = [f for f in entity.fields if f != external_field]
    lookup_targets = set(entity.lookups)
    rng = random.Random(f"{seed}:{entity.name}")  # nosec B311 -- deterministic audit sampling, not security
    sampled = rng.sample(sorted(source_by_key), min(RECONCILE_SAMPLE, len(source_by_key)))
    for key in sampled:
        source_row = source_by_key[key]
        target_row = target_by_key.get(key)
        if target_row is None:
            continue  # Already counted as missing.
        for target in comparable:
            if target in lookup_targets or target not in target_header:
                continue
            expected, failed = mapped_value(entity, target, source_row)
            if failed is not None:
                continue
            actual = target_row.get(target, "")
            if (expected or "") != (actual or ""):
                recap.sample_diffs.append((key, target, expected or "", actual or ""))
    # Relationships: source children whose lookup parent key is absent.
    # Lookups pointing outside the mapped entities cannot be judged: skip.
    for target, lookup in entity.lookups.items():
        _ = target
        if lookup.entity not in parent_keys:
            continue
        via = lookup.via
        known_parents = parent_keys[lookup.entity]
        for row in source_rows:
            parent_key = (row.get(via) or "").strip()
            if not parent_key or parent_key in known_parents:
                continue
            child_key = _external_key(entity, row) or dedupe_mod.row_key(row, None)
            recap.relationship_gaps.append((child_key, via, parent_key))
    # Stage distribution (source stage column vs target Stage field).
    stage_from: str | None = None
    for target, field_map in entity.fields.items():
        if target == "Stage" and field_map.from_col is not None:
            stage_from = field_map.from_col
    if stage_from is not None and "Stage" in target_header:
        source_dist: dict[str, int] = {}
        for row in source_rows:
            value, failed = mapped_value(entity, "Stage", row)
            if failed is None and value:
                source_dist[value] = source_dist.get(value, 0) + 1
        target_dist: dict[str, int] = {}
        for row in target_rows:
            value = (row.get("Stage") or "").strip()
            if value:
                target_dist[value] = target_dist.get(value, 0) + 1
        for stage in sorted(set(source_dist) | set(target_dist)):
            recap.stage_rows.append((stage, source_dist.get(stage, 0), target_dist.get(stage, 0)))
    return recap


def recap_findings(recaps: list[EntityRecap]) -> list[Finding]:
    """Findings for count gaps, sample diffs and relationship orphans."""
    findings: list[Finding] = []
    for recap in recaps:
        if recap.missing_keys or recap.extra_keys:
            findings.append(
                Finding.create(
                    module="migration",
                    code="reconcile_count_mismatch",
                    severity=Severity.ERROR,
                    entity=recap.entity,
                    entity_id="counts",
                    message="Source and target counts differ for the module.",
                    evidence={
                        "source_total": recap.source_total,
                        "target_total": recap.target_total,
                        "missing": len(recap.missing_keys),
                        "extra": len(recap.extra_keys),
                    },
                    remediation="Import the missing rows or explain the extras.",
                    discriminator="reconcile\0counts",
                )
            )
        seen_rows: set[str] = set()
        for key, target, _, _ in recap.sample_diffs:
            if key in seen_rows:
                continue
            seen_rows.add(key)
            findings.append(
                Finding.create(
                    module="migration",
                    code="reconcile_field_diff",
                    severity=Severity.REVIEW,
                    entity=recap.entity,
                    entity_id=key,
                    message="Sampled target record differs from the transformed source.",
                    evidence={"target_field": target},
                    remediation="Open the Sample diffs tab and fix the import or the source.",
                    discriminator=f"reconcile\0sample\0{key}",
                )
            )
        for child_key, via, _ in recap.relationship_gaps:
            findings.append(
                Finding.create(
                    module="migration",
                    code="reconcile_relationship_gap",
                    severity=Severity.ERROR,
                    entity=recap.entity,
                    entity_id=child_key,
                    message="Child row references a parent absent from the source.",
                    evidence={"via": via},
                    remediation="Import the parent first or fix the lookup value.",
                    discriminator=f"reconcile\0relationship\0{child_key}\0{via}",
                )
            )
    return findings


def _sheet_header(sheet: Worksheet, titles: list[str], *, row: int = 1) -> None:
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = _BOLD


def render_reconcile_workbook(recaps: list[EntityRecap], *, seed: int) -> bytes:
    """Render the four-tab reconcile workbook (seed recorded on Summary)."""
    book = Workbook()
    summary = book.worksheets[0]
    summary.title = "Summary"
    summary["A1"] = "Migration reconciliation"
    summary["A1"].font = _BOLD
    summary["A2"] = "Sample seed"
    summary["B2"] = seed
    _sheet_header(summary, ["Entity", "Module", "Source", "Target", "Missing", "Extra"], row=4)
    for number, recap in enumerate(recaps, start=5):
        summary.cell(row=number, column=1, value=recap.entity)
        summary.cell(row=number, column=2, value=recap.module)
        summary.cell(row=number, column=3, value=recap.source_total)
        summary.cell(row=number, column=4, value=recap.target_total)
        summary.cell(row=number, column=5, value=len(recap.missing_keys))
        summary.cell(row=number, column=6, value=len(recap.extra_keys))
    summary.freeze_panes = "A5"
    counts = book.create_sheet("Counts")
    _sheet_header(counts, ["Entity", "Stage", "Source", "Target", "Delta"])
    line = 2
    for recap in recaps:
        for stage, source_count, target_count in recap.stage_rows:
            counts.cell(row=line, column=1, value=recap.entity)
            counts.cell(row=line, column=2, value=stage)
            counts.cell(row=line, column=3, value=source_count)
            counts.cell(row=line, column=4, value=target_count)
            counts.cell(row=line, column=5, value=f"=C{line}-D{line}")
            line += 1
    if line == 2:
        counts.cell(row=2, column=1, value="no stage distribution available")
    counts.freeze_panes = "A2"
    diffs = book.create_sheet("Sample diffs")
    _sheet_header(diffs, ["Entity", "Key", "Field", "Source", "Target", "Match"])
    line = 2
    for recap in recaps:
        for key, target, expected, actual in recap.sample_diffs:
            diffs.cell(row=line, column=1, value=recap.entity)
            diffs.cell(row=line, column=2, value=key)
            diffs.cell(row=line, column=3, value=target)
            diffs.cell(row=line, column=4, value=expected)
            diffs.cell(row=line, column=5, value=actual)
            diffs.cell(row=line, column=6, value="no")
            line += 1
    if line == 2:
        diffs.cell(row=2, column=1, value="no sample diffs")
    diffs.freeze_panes = "A2"
    relations = book.create_sheet("Relationships")
    _sheet_header(relations, ["Entity", "Child", "Via", "Parent", "Side"])
    line = 2
    for recap in recaps:
        for child_key, via, parent_key in recap.relationship_gaps:
            relations.cell(row=line, column=1, value=recap.entity)
            relations.cell(row=line, column=2, value=child_key)
            relations.cell(row=line, column=3, value=via)
            relations.cell(row=line, column=4, value=parent_key)
            relations.cell(row=line, column=5, value="source")
            line += 1
    if line == 2:
        relations.cell(row=2, column=1, value="no relationship gaps")
    relations.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def reconcile(
    doc: MappingDoc,
    source_paths: dict[str, str],
    target_dir: str,
    *,
    seed: int = 42,
) -> tuple[list[EntityRecap], bytes]:
    """Reconcile mapped sources against target dumps; return recaps + workbook."""
    from pathlib import Path

    parent_keys: dict[str, set[str]] = {}
    for other in doc.entities:
        other_path = source_paths.get(other.name)
        if other_path is None or other.external_id is None:
            continue
        parent_keys[other.name] = {
            key
            for row in _read_csv_rows(other_path)[1]
            if (key := _external_key(other, row)) is not None
        }
    recaps: list[EntityRecap] = []
    for entity in doc.entities:
        source_path = source_paths.get(entity.name)
        if source_path is None:
            continue
        _, source_rows = _read_csv_rows(source_path)
        target_path = Path(target_dir) / f"{entity.target_module}.csv"
        if not target_path.is_file():
            recap = EntityRecap(entity=entity.name, module=entity.target_module)
            recap.source_total = len(source_rows)
            recap.notes.append(f"target file {target_path.name} absent: target counts unknown")
            recaps.append(recap)
            continue
        target_header, target_rows = _read_csv_rows(str(target_path))
        recaps.append(
            reconcile_entity(
                entity, source_rows, target_header, target_rows, parent_keys, seed=seed
            )
        )
    return recaps, render_reconcile_workbook(recaps, seed=seed)


__all__: list[str] = [
    "RECONCILE_SAMPLE",
    "RECONCILE_SHEETS",
    "EntityRecap",
    "recap_findings",
    "reconcile",
    "reconcile_entity",
    "render_reconcile_workbook",
]
