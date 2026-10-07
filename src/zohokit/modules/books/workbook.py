"""Finance sign-off workbook for Books month-end recon (TK-BK-F7).

Seven tabs: ``Matched`` / ``Mismatched`` / ``Missing invoice`` /
``Orphan invoice`` / ``Needs review`` / ``FX`` / ``Sign-off``. Amount
columns carry real per-currency ``=SUMIF(...)`` totals formulas (finance
can audit the arithmetic in Excel — one ``Total INR`` / ``Total USD``
row per currency, so mixed-currency tabs never add INR and USD
together); converted amounts total only in their converted currency.
Amount cells use the ``#,##0.00`` money format. Every other cell is
literal text through :func:`safe_text`, so data-sourced values (deal
names, customers, reference numbers, messages) can never become
formulas (STD §4 safety — only the toolkit's own totals cells may be
formulas).
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from zohokit.core.findings import Report
from zohokit.reports.xlsx import safe_text

#: Workbook tabs in order.
BOOKS_SHEETS = (
    "Matched",
    "Mismatched",
    "Missing invoice",
    "Orphan invoice",
    "Needs review",
    "FX",
    "Sign-off",
)

#: Row columns for the five finding tabs. ``Record`` is the finding's
#: record type (deal/invoice/credit_note/entity); ``Legal entity`` is the
#: entity key (``in-entity``/``us-entity``) from the finding evidence.
ROW_COLUMNS = (
    "Record",
    "Legal entity",
    "ID",
    "Customer",
    "Code",
    "Deal Net",
    "Invoiced Total",
    "Currency",
    "Converted Total",
    "Converted Currency",
    "Message",
)

#: Row columns for the FX tab (one row per conversion or missing rate).
FX_COLUMNS = (
    "Deal",
    "Invoice",
    "Original",
    "Original Currency",
    "Converted",
    "Converted Currency",
    "Rate",
    "Rate Source",
    "Note",
)

_MATCHED_CODES = frozenset({"within_tolerance"})
_MISMATCHED_CODES = frozenset({"outside_tolerance", "under_invoiced", "over_invoiced"})

_BOLD = Font(bold=True)

#: Excel number format for amount cells (finance-readable money).
MONEY_FORMAT = "#,##0.00"


def _tab_for(code: str) -> str:
    """Finding tab for *code* (anything else needs review)."""
    if code in _MATCHED_CODES:
        return "Matched"
    if code in _MISMATCHED_CODES:
        return "Mismatched"
    if code == "missing_invoice":
        return "Missing invoice"
    if code == "orphan_invoice_reference":
        return "Orphan invoice"
    return "Needs review"


def _number(raw: Any) -> float | None:
    """Evidence amount as a float for Excel, else ``None`` (blank cell)."""
    if raw is None or raw == "":
        return None
    try:
        return float(str(raw))
    except ValueError:
        return None


def _legal_entity(finding: dict[str, Any]) -> str:
    """Legal-entity key for a finding row (never a fingerprint or org id)."""
    evidence = finding.get("evidence") or {}
    record = str(finding.get("entity", ""))
    preferred = {
        "deal": "deal_entity",
        "invoice": "invoice_entity",
        "credit_note": "credit_note_entity",
        "entity": "entity_key",
    }.get(record, "")
    if preferred and evidence.get(preferred):
        return str(evidence[preferred])
    for key in ("deal_entity", "invoice_entity", "credit_note_entity", "entity_key"):
        if evidence.get(key):
            return str(evidence[key])
    return ""


def _row_for(finding: dict[str, Any]) -> list[Any]:
    """One workbook row for a recon finding.

    Needs-review rows show the known amounts too (e.g. ``fx_rate_missing``
    carries the deal net plus the original invoice amount and currency):
    ``Invoiced Total`` falls back to the invoice/credit-note amount and
    ``Currency`` to its currency when the deal currency is absent.
    """
    evidence = finding.get("evidence") or {}
    conversions = evidence.get("conversions") or []
    converted_total: Any = ""
    converted_ccy = ""
    if isinstance(conversions, list) and conversions:
        first = conversions[0] if isinstance(conversions[0], dict) else {}
        parsed = _number(first.get("converted"))
        converted_total = parsed if parsed is not None else ""
        converted_ccy = str(first.get("converted_currency", ""))
    invoiced_raw = evidence.get("invoiced_total")
    if invoiced_raw in (None, ""):
        invoiced_raw = evidence.get("invoice_amount")
    if invoiced_raw in (None, ""):
        invoiced_raw = evidence.get("credit_note_amount")
    currency = str(evidence.get("currency", ""))
    if not currency:
        currency = str(evidence.get("invoice_currency", ""))
    if not currency:
        currency = str(evidence.get("credit_note_currency", ""))
    return [
        str(finding.get("entity", "")),
        _legal_entity(finding),
        str(finding.get("entity_id", "")),
        str(evidence.get("customer_name", "")),
        str(finding.get("code", "")),
        _number(evidence.get("deal_net")),
        _number(invoiced_raw),
        currency,
        converted_total,
        converted_ccy,
        str(finding.get("message", "")),
    ]


def _fx_rows(findings: list[dict[str, Any]]) -> list[list[Any]]:
    """FX tab rows: one per conversion plus one per missing rate."""
    rows: list[list[Any]] = []
    for finding in findings:
        evidence = finding.get("evidence") or {}
        conversions = evidence.get("conversions")
        if isinstance(conversions, list):
            for entry in conversions:
                if not isinstance(entry, dict):
                    continue
                rows.append(
                    [
                        str(finding.get("entity_id", "")),
                        str(entry.get("invoice_id", "")),
                        _number(entry.get("original")),
                        str(entry.get("original_currency", "")),
                        _number(entry.get("converted")),
                        str(entry.get("converted_currency", "")),
                        str(entry.get("rate", "")),
                        str(entry.get("source", entry.get("rate_source", ""))),
                        "",
                    ]
                )
        if finding.get("code") == "fx_rate_missing":
            original = _number(evidence.get("invoice_amount"))
            rows.append(
                [
                    str(finding.get("entity_id", "")),
                    str(evidence.get("invoice_id", "")),
                    original if original is not None else "",
                    str(evidence.get("invoice_currency", "") or evidence.get("from_currency", "")),
                    "",
                    str(evidence.get("to_currency", "")),
                    "",
                    "",
                    "no exact rate: comparison skipped, never guessed",
                ]
            )
    return rows


def _write_header(sheet: Worksheet, columns: tuple[str, ...]) -> None:
    for number, title in enumerate(columns, start=1):
        sheet.cell(row=1, column=number, value=title).font = _BOLD


def _write_rows(
    sheet: Worksheet, rows: list[list[Any]], amount_columns: tuple[int, ...] = ()
) -> int:
    """Write data rows through :func:`safe_text`; return the totals row.

    Numeric cells in *amount_columns* get the ``#,##0.00`` money format;
    every other cell stays literal text.
    """
    money_cols = set(amount_columns)
    for line, values in enumerate(rows, start=2):
        for number, value in enumerate(values, start=1):
            cell = sheet.cell(row=line, column=number)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.value = value
                if number in money_cols:
                    cell.number_format = MONEY_FORMAT
            else:
                cell.value = safe_text(value)
    return max(len(rows) + 2, 3)


def _totals_block(
    sheet: Worksheet,
    first_row: int,
    last_data: int,
    groups: Sequence[tuple[str, int, str, Sequence[int]]],
    *,
    fallback_sums: Sequence[int],
) -> None:
    """Write one totals row per currency group, or a plain SUM row.

    Each group is ``(label, criteria_column, currency, sum_columns)``
    rendered as ``=SUMIF(C2:C7,"INR",F2:F7)``. With no groups (an empty
    tab) a single ``Total`` row carries plain ``=SUM`` formulas that
    recompute to zero.
    """

    def _letter(column: int) -> str:
        return str(sheet.cell(row=1, column=column).column_letter)

    if not groups:
        sheet.cell(row=first_row, column=1, value="Total").font = _BOLD
        for column in fallback_sums:
            letter = _letter(column)
            sheet.cell(row=first_row, column=column, value=f"=SUM({letter}2:{letter}{last_data})")
        return
    row = first_row
    for label, criteria_column, currency, sum_columns in groups:
        sheet.cell(row=row, column=1, value=label).font = _BOLD
        criteria_letter = _letter(criteria_column)
        for column in sum_columns:
            sum_letter = _letter(column)
            sheet.cell(
                row=row,
                column=column,
                value=(
                    f"=SUMIF({criteria_letter}2:{criteria_letter}{last_data},"
                    f'"{currency}",{sum_letter}2:{sum_letter}{last_data})'
                ),
            )
        row += 1


def _write_finding_totals(sheet: Worksheet, rows: list[list[Any]]) -> None:
    """Per-currency totals for a finding tab (never mixing currencies)."""
    last_data = max(len(rows) + 1, 2)
    first_row = max(len(rows) + 2, 3)
    deal_ccys = sorted({str(values[7]) for values in rows if str(values[7])})
    converted_ccys = sorted({str(values[9]) for values in rows if str(values[9])})
    groups: list[tuple[str, int, str, Sequence[int]]] = [
        (f"Total {ccy}", 8, ccy, (6, 7)) for ccy in deal_ccys
    ]
    groups += [(f"Total converted {ccy}", 10, ccy, (9,)) for ccy in converted_ccys]
    _totals_block(sheet, first_row, last_data, groups, fallback_sums=(6, 7))


def _write_fx_totals(sheet: Worksheet, rows: list[list[Any]]) -> None:
    """Per-currency totals for the FX tab (original + converted)."""
    last_data = max(len(rows) + 1, 2)
    first_row = max(len(rows) + 2, 3)
    original_ccys = sorted({str(values[3]) for values in rows if str(values[3])})
    converted_ccys = sorted({str(values[5]) for values in rows if str(values[5])})
    groups: list[tuple[str, int, str, Sequence[int]]] = [
        (f"Total original {ccy}", 4, ccy, (3,)) for ccy in original_ccys
    ]
    groups += [(f"Total converted {ccy}", 6, ccy, (5,)) for ccy in converted_ccys]
    _totals_block(sheet, first_row, last_data, groups, fallback_sums=(3, 5))


def _finish_tab(sheet: Worksheet, columns: tuple[str, ...]) -> None:
    sheet.freeze_panes = "A2"
    last_col = sheet.cell(row=1, column=len(columns)).column_letter
    sheet.auto_filter.ref = f"A1:{last_col}{sheet.max_row}"


def render_books_workbook(report: Report) -> bytes:
    """Render a recon *report* as the finance sign-off workbook."""
    dumped = report.model_dump(mode="json")
    findings = dumped.get("findings", [])
    book = Workbook()
    book.worksheets[0].title = BOOKS_SHEETS[0]
    sheets = {
        name: (book.worksheets[0] if name == BOOKS_SHEETS[0] else book.create_sheet(name))
        for name in BOOKS_SHEETS
    }
    buckets: dict[str, list[dict[str, Any]]] = {
        "Matched": [],
        "Mismatched": [],
        "Missing invoice": [],
        "Orphan invoice": [],
        "Needs review": [],
    }
    for finding in findings:
        buckets[_tab_for(str(finding.get("code", "")))].append(finding)
    for name in ("Matched", "Mismatched", "Missing invoice", "Orphan invoice", "Needs review"):
        sheet = sheets[name]
        _write_header(sheet, ROW_COLUMNS)
        rows = [_row_for(item) for item in buckets[name]]
        _write_rows(sheet, rows, amount_columns=(6, 7, 9))
        _write_finding_totals(sheet, rows)
        _finish_tab(sheet, ROW_COLUMNS)
    fx_sheet = sheets["FX"]
    _write_header(fx_sheet, FX_COLUMNS)
    fx_data = _fx_rows(findings)
    _write_rows(fx_sheet, fx_data, amount_columns=(3, 5))
    _write_fx_totals(fx_sheet, fx_data)
    _finish_tab(fx_sheet, FX_COLUMNS)
    signoff = sheets["Sign-off"]
    signoff["A1"] = "Sign-off"
    signoff["A1"].font = _BOLD
    for number, label in enumerate(("Preparer", "Reviewer", "Date"), start=3):
        signoff.cell(row=number, column=1, value=label).font = _BOLD
        signoff.cell(row=number, column=2, value="")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def sum_range(formula: str, sheet: Worksheet) -> float:
    """Evaluate a ``=SUM(Xa:Xb)`` *formula* against *sheet* values (tests)."""
    match = re.fullmatch(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", formula.strip())
    if match is None:
        raise ValueError(f"not a SUM range formula: {formula!r}")
    col_a, row_a, col_b, row_b = (
        match.group(1),
        int(match.group(2)),
        match.group(3),
        int(match.group(4)),
    )
    if col_a != col_b:
        raise ValueError(f"not a single-column SUM: {formula!r}")
    total = 0.0
    for row in range(row_a, row_b + 1):
        value = sheet[f"{col_a}{row}"].value
        if isinstance(value, bool) or value is None:
            continue
        if isinstance(value, (int, float)):
            total += float(value)
    return total


def sum_if_range(formula: str, sheet: Worksheet) -> float:
    """Evaluate a ``=SUMIF(Ga:Gb,"CCY",Ea:Eb)`` *formula* (tests)."""
    match = re.fullmatch(
        r'=SUMIF\(([A-Z]+)(\d+):([A-Z]+)(\d+),"([^"]+)",([A-Z]+)(\d+):([A-Z]+)(\d+)\)',
        formula.strip(),
    )
    if match is None:
        raise ValueError(f"not a SUMIF range formula: {formula!r}")
    crit_col, crit_first, _, crit_last, currency, sum_col, sum_first, _, sum_last = (
        match.group(1),
        int(match.group(2)),
        match.group(3),
        int(match.group(4)),
        match.group(5),
        match.group(6),
        int(match.group(7)),
        match.group(8),
        int(match.group(9)),
    )
    if (crit_first, crit_last) != (sum_first, sum_last):
        raise ValueError(f"SUMIF ranges must align: {formula!r}")
    total = 0.0
    for row in range(crit_first, crit_last + 1):
        if str(sheet[f"{crit_col}{row}"].value) != currency:
            continue
        value = sheet[f"{sum_col}{row}"].value
        if isinstance(value, bool) or value is None:
            continue
        if isinstance(value, (int, float)):
            total += float(value)
    return total


__all__: list[str] = [
    "BOOKS_SHEETS",
    "FX_COLUMNS",
    "MONEY_FORMAT",
    "ROW_COLUMNS",
    "render_books_workbook",
    "sum_if_range",
    "sum_range",
]
